from __future__ import annotations

import io
import mimetypes
import os
import re
import sys
import time
import urllib.error
import urllib.request
import warnings
import zipfile
from html import unescape
from typing import Any
from urllib.parse import unquote, urlparse
from xml.etree import ElementTree as ET

from braintrust import Attachment
from openai import BadRequestError, OpenAI, RateLimitError

DEFAULT_MODEL = "gpt-5.4-mini"
DEFAULT_JUDGE_MODEL = "gpt-5.4-mini"
DEFAULT_JUDGE_MAX_TOKENS = 64

REFERENCE_DOCUMENT_SUFFIXES = frozenset({".xlsx", ".docx", ".pptx", ".pdf"})
GENERATABLE_SUFFIXES = frozenset({".xlsx", ".docx", ".pptx", ".pdf"})
MAX_REFERENCE_DOWNLOAD_BYTES = 25 * 1024 * 1024
MAX_EXTRACT_CHARS_PER_FILE = 250_000
MAX_REFERENCE_CONTENT_TOTAL_CHARS = 500_000
MAX_GENERATED_FILES = 4
MAX_GENERATED_CONTENT_CHARS_PER_FILE = 80_000
FETCH_TIMEOUT_SEC = 90
MAX_OPENAI_RETRIES = 12
INITIAL_OPENAI_BACKOFF_SEC = 1.0


def suffix_from_url(url: str) -> str:
    path = unquote(urlparse(url).path)
    return os.path.splitext(path)[1].lower()


def fetch_url_bytes(url: str, *, user_agent: str) -> bytes | None:
    req = urllib.request.Request(url, headers={"User-Agent": user_agent})
    try:
        with urllib.request.urlopen(req, timeout=FETCH_TIMEOUT_SEC) as resp:
            chunks: list[bytes] = []
            total = 0
            while total < MAX_REFERENCE_DOWNLOAD_BYTES:
                chunk = resp.read(min(65536, MAX_REFERENCE_DOWNLOAD_BYTES - total))
                if not chunk:
                    break
                chunks.append(chunk)
                total += len(chunk)
            if resp.read(1):
                print(f"WARNING: Skipping oversized reference file: {url}", file=sys.stderr)
                return None
            return b"".join(chunks)
    except (urllib.error.URLError, TimeoutError, OSError) as err:
        print(f"WARNING: Could not download reference file {url}: {err}", file=sys.stderr)
        return None


def truncate_text(text: str, limit: int) -> str:
    if len(text) <= limit:
        return text
    return text[:limit] + f"\n\n[... truncated {len(text) - limit} characters ...]"


def guess_content_type(path: str) -> str:
    guessed, _ = mimetypes.guess_type(path)
    return guessed or "application/octet-stream"


def make_file_attachment(path: str, filename: str | None = None) -> Attachment:
    return Attachment(
        data=path,
        filename=filename or os.path.basename(path),
        content_type=guess_content_type(path),
    )


def make_bytes_attachment(data: bytes, filename: str) -> Attachment:
    return Attachment(
        data=data,
        filename=filename,
        content_type=guess_content_type(filename),
    )


def source_file_attachments(source_file_paths: tuple[str, ...]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for path in source_file_paths:
        if not os.path.exists(path):
            continue
        out.append({"path": path, "attachment": make_file_attachment(path)})
    return out


def extract_xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook

    buf = io.BytesIO(data)
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", category=UserWarning, module=r"openpyxl\..*")
        wb = load_workbook(buf, read_only=True, data_only=True)
        try:
            parts: list[str] = []
            for sheet in wb:
                rows_out: list[str] = []
                for row in sheet.iter_rows(values_only=True):
                    cells = ["" if c is None else str(c) for c in row]
                    if any(cells):
                        rows_out.append("\t".join(cells))
                if rows_out:
                    parts.append(f"## Sheet: {sheet.title}\n" + "\n".join(rows_out))
            return "\n\n".join(parts) if parts else "(empty workbook)"
        finally:
            wb.close()


def extract_docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    return "\n".join(parts) if parts else "(empty document)"


def extract_docx_text_fallback(data: bytes) -> str:
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

    def part_text(xml_bytes: bytes) -> list[str]:
        try:
            root = ET.fromstring(xml_bytes)
            lines: list[str] = []
            for paragraph in root.findall(".//w:p", namespace):
                fragments: list[str] = []
                for node in paragraph.iter():
                    tag = node.tag.rsplit("}", 1)[-1] if "}" in node.tag else node.tag
                    if tag == "t" and node.text:
                        fragments.append(node.text)
                    elif tag == "tab":
                        fragments.append("\t")
                    elif tag in {"br", "cr"}:
                        fragments.append("\n")
                text = "".join(fragments).strip()
                if text:
                    lines.append(text)
            return lines
        except ET.ParseError:
            text = re.sub(r"<[^>]+>", " ", xml_bytes.decode("utf-8", errors="ignore"))
            text = re.sub(r"\s+", " ", unescape(text)).strip()
            return [text] if text else []

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            collected: list[str] = []
            for name in zf.namelist():
                if not name.startswith("word/"):
                    continue
                if not name.endswith(".xml"):
                    continue
                if "document.xml" not in name and "header" not in name and "footer" not in name:
                    continue
                collected.extend(part_text(zf.read(name)))
    except (zipfile.BadZipFile, KeyError, OSError):
        return ""

    return "\n".join(collected) if collected else ""


def extract_pptx_text(data: bytes) -> str:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(io.BytesIO(data))

    def shapes_text(shapes) -> list[str]:
        lines: list[str] = []
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                lines.extend(shapes_text(shape.shapes))
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    t = "".join(run.text for run in paragraph.runs).strip()
                    if t:
                        lines.append(t)
        return lines

    parts: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        slide_lines = shapes_text(slide.shapes)
        if slide_lines:
            parts.append(f"--- Slide {i} ---\n" + "\n".join(slide_lines))
    return "\n\n".join(parts) if parts else "(empty presentation)"


def extract_pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page_num, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            parts.append(f"--- Page {page_num} ---\n{text}")
    return "\n\n".join(parts) if parts else "(empty pdf)"


def extract_reference_document_text(source: str, data: bytes, suffix: str) -> str:
    try:
        if suffix == ".xlsx":
            raw = extract_xlsx_text(data)
        elif suffix == ".docx":
            try:
                raw = extract_docx_text(data)
            except Exception:
                raw = extract_docx_text_fallback(data)
        elif suffix == ".pptx":
            raw = extract_pptx_text(data)
        elif suffix == ".pdf":
            raw = extract_pdf_text(data)
        else:
            return ""
    except Exception as err:
        print(f"WARNING: Could not parse document {source}: {err}", file=sys.stderr)
        return ""
    return truncate_text(raw, MAX_EXTRACT_CHARS_PER_FILE)


def build_reference_file_context(
    urls: list[str],
    *,
    user_agent: str,
) -> tuple[str, list[dict[str, Any]]]:
    blocks: list[str] = []
    attachments: list[dict[str, Any]] = []
    total_budget = MAX_REFERENCE_CONTENT_TOTAL_CHARS

    for url in urls[:50]:
        suffix = suffix_from_url(url)
        if suffix not in REFERENCE_DOCUMENT_SUFFIXES:
            continue
        data = fetch_url_bytes(url, user_agent=user_agent)
        if not data:
            continue
        body = extract_reference_document_text(url, data, suffix)
        if not body:
            continue
        attachments.append(
            {
                "url": url,
                "attachment": make_bytes_attachment(data, os.path.basename(unquote(urlparse(url).path)) or f"reference{suffix}"),
            }
        )
        header = f"### Reference file ({suffix}): {url}\n"
        room = total_budget - len(header)
        if room <= 0:
            blocks.append(
                "### Reference files\n"
                f"[Omitted remaining reference files: budget {MAX_REFERENCE_CONTENT_TOTAL_CHARS} chars exceeded]"
            )
            break
        if len(body) > room:
            body = truncate_text(body, room)
        blocks.append(header + body)
        total_budget -= len(header) + len(body)

    if not blocks:
        return "", attachments
    return "\n\n".join(["## Extracted reference file contents", *blocks]), attachments


def format_reference_files_section(
    urls: list[str],
    *,
    user_agent: str,
) -> tuple[str, list[dict[str, Any]]]:
    if not urls:
        return "", []
    url_list = "\n".join(f"  - {u}" for u in urls[:50])
    doc_blob, attachments = build_reference_file_context(urls, user_agent=user_agent)
    parts = [f"\n\nReference files:\n{url_list}"]
    if doc_blob:
        parts.append("\n\n" + doc_blob)
    return "".join(parts), attachments


def sanitize_filename(filename: str, fallback_suffix: str) -> str:
    candidate = (filename or "").strip()
    candidate = candidate.replace("\\", "_").replace("/", "_")
    candidate = re.sub(r"[^A-Za-z0-9._ -]+", "_", candidate).strip(" .")
    if not candidate:
        candidate = f"deliverable{fallback_suffix}"
    suffix = os.path.splitext(candidate)[1].lower()
    if suffix in {".xls", ".xlsm"}:
        candidate = os.path.splitext(candidate)[0] + ".xlsx"
    elif suffix == ".ppt":
        candidate = os.path.splitext(candidate)[0] + ".pptx"
    elif suffix not in GENERATABLE_SUFFIXES:
        candidate = os.path.splitext(candidate)[0] + fallback_suffix
    return candidate


def sanitize_excel_sheet_title(title: str, fallback: str) -> str:
    cleaned = re.sub(r"[:\\/*?\[\]]", "_", (title or "").strip())
    cleaned = cleaned.strip("'")
    cleaned = cleaned[:31].strip()
    return cleaned or fallback[:31] or "Sheet"


def string_rows(raw_rows: Any) -> list[list[str]]:
    if not isinstance(raw_rows, list):
        return []
    rows: list[list[str]] = []
    for row in raw_rows:
        if isinstance(row, list):
            rows.append([str(cell) if cell is not None else "" for cell in row])
    return rows


def build_docx(path: str, spec: dict) -> None:
    from docx import Document

    doc = Document()
    title = str(spec.get("title") or "").strip()
    if title:
        doc.add_heading(title, level=1)

    paragraphs = spec.get("paragraphs")
    if isinstance(paragraphs, list):
        for paragraph in paragraphs:
            text = str(paragraph).strip()
            if text:
                doc.add_paragraph(text)
    else:
        body = str(spec.get("content") or spec.get("text") or "").strip()
        if body:
            for paragraph in body.splitlines():
                if paragraph.strip():
                    doc.add_paragraph(paragraph.strip())

    tables = spec.get("tables")
    if isinstance(tables, list):
        for table_spec in tables:
            rows = string_rows((table_spec or {}).get("rows"))
            if not rows:
                continue
            tbl = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            for r_idx, row in enumerate(rows):
                for c_idx, value in enumerate(row):
                    tbl.cell(r_idx, c_idx).text = value

    if len(doc.paragraphs) == 0 and len(doc.tables) == 0:
        doc.add_paragraph("(empty deliverable)")
    doc.save(path)


def build_xlsx(path: str, spec: dict) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)
    used_names: set[str] = set()

    sheets = spec.get("sheets")
    if isinstance(sheets, list):
        for idx, sheet_spec in enumerate(sheets, start=1):
            base_name = sanitize_excel_sheet_title(str((sheet_spec or {}).get("name") or f"Sheet{idx}"), f"Sheet{idx}")
            name = base_name
            counter = 2
            while name in used_names:
                suffix = f"_{counter}"
                name = (base_name[: max(0, 31 - len(suffix))] + suffix) or f"Sheet{idx}"
                counter += 1
            used_names.add(name)
            ws = wb.create_sheet(title=name)
            for row in string_rows((sheet_spec or {}).get("rows")):
                ws.append(row)
    else:
        ws = wb.create_sheet(title="Sheet1")
        raw = str(spec.get("content") or spec.get("text") or "").strip()
        if raw:
            for line in raw.splitlines():
                ws.append([cell.strip() for cell in line.split("\t")])

    if len(wb.sheetnames) == 0:
        ws = wb.create_sheet(title="Sheet1")
        ws["A1"] = "(empty deliverable)"
    wb.save(path)
    wb.close()


def build_pptx(path: str, spec: dict) -> None:
    from pptx import Presentation

    prs = Presentation()
    slides = spec.get("slides")
    if isinstance(slides, list) and slides:
        for slide_spec in slides:
            layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(layout)
            title = str((slide_spec or {}).get("title") or "").strip()
            body_lines = slide_spec.get("bullets")
            if not isinstance(body_lines, list):
                fallback = str((slide_spec or {}).get("content") or "").strip()
                body_lines = [line.strip() for line in fallback.splitlines() if line.strip()]
            slide.shapes.title.text = title or "Slide"
            text_frame = slide.shapes.placeholders[1].text_frame
            text_frame.clear()
            if body_lines:
                text_frame.text = str(body_lines[0])
                for line in body_lines[1:]:
                    p = text_frame.add_paragraph()
                    p.text = str(line)
            else:
                text_frame.text = "(empty)"
    else:
        layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = str(spec.get("title") or "Slide 1")
        text_frame = slide.shapes.placeholders[1].text_frame
        text_frame.text = str(spec.get("content") or spec.get("text") or "(empty)")

    prs.save(path)


def build_pdf(path: str, spec: dict) -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    lines: list[str] = []
    title = str(spec.get("title") or "").strip()
    if title:
        lines.append(title)
        lines.append("")
    paragraphs = spec.get("paragraphs")
    if isinstance(paragraphs, list):
        for paragraph in paragraphs:
            lines.extend(str(paragraph).splitlines())
            lines.append("")
    else:
        body = str(spec.get("content") or spec.get("text") or "").strip()
        lines.extend(body.splitlines() or ["(empty deliverable)"])

    c = canvas.Canvas(path, pagesize=letter)
    _, height = letter
    y = height - 50
    for raw_line in lines:
        line = raw_line.rstrip()
        chunks = [line[i : i + 110] for i in range(0, len(line), 110)] or [""]
        for chunk in chunks:
            c.drawString(50, y, chunk)
            y -= 14
            if y < 50:
                c.showPage()
                y = height - 50
    c.save()


def extract_generated_content(path: str, suffix: str) -> str:
    if suffix in REFERENCE_DOCUMENT_SUFFIXES:
        with open(path, "rb") as f:
            data = f.read()
        return extract_reference_document_text(path, data, suffix)
    return ""


def build_artifacts(
    task_id: str,
    payload: dict[str, Any],
    allowed_suffixes: list[str],
    *,
    artifact_output_dir: str,
) -> list[dict[str, Any]]:
    task_dir = os.path.join(artifact_output_dir, task_id)
    os.makedirs(task_dir, exist_ok=True)

    artifacts = payload.get("artifacts")
    if not isinstance(artifacts, list):
        return []

    out: list[dict[str, Any]] = []
    for artifact in artifacts[:MAX_GENERATED_FILES]:
        if not isinstance(artifact, dict):
            continue
        requested_type = str(artifact.get("file_type") or "").strip().lower()
        filename = str(artifact.get("filename") or "").strip()
        suffix = os.path.splitext(filename)[1].lower()
        if requested_type and not requested_type.startswith("."):
            requested_type = "." + requested_type
        if requested_type in {".xls", ".xlsm"}:
            requested_type = ".xlsx"
        if requested_type == ".ppt":
            requested_type = ".pptx"
        if suffix in {".xls", ".xlsm"}:
            suffix = ".xlsx"
        elif suffix == ".ppt":
            suffix = ".pptx"
        target_suffix = suffix or requested_type
        if target_suffix not in GENERATABLE_SUFFIXES:
            target_suffix = allowed_suffixes[0] if allowed_suffixes else ".docx"
        if target_suffix not in allowed_suffixes:
            continue

        safe_name = sanitize_filename(filename or f"deliverable{target_suffix}", target_suffix)
        path = os.path.join(task_dir, safe_name)

        try:
            if target_suffix == ".docx":
                build_docx(path, artifact)
            elif target_suffix == ".xlsx":
                build_xlsx(path, artifact)
            elif target_suffix == ".pptx":
                build_pptx(path, artifact)
            elif target_suffix == ".pdf":
                build_pdf(path, artifact)
            else:
                continue
            out.append(
                {
                    "filename": safe_name,
                    "suffix": target_suffix,
                    "path": path,
                    "content": extract_generated_content(path, target_suffix),
                }
            )
        except Exception as err:
            print(f"WARNING: Could not create artifact {safe_name}: {err}", file=sys.stderr)
    return out


def format_artifact_output(narrative: str, artifacts: list[dict[str, Any]]) -> str:
    parts = [f"NARRATIVE RESPONSE:\n{(narrative or '').strip()}"]
    if not artifacts:
        parts.append("\nGENERATED FILES:\n(none)")
        return "\n\n".join(parts)

    lines = ["\nGENERATED FILES:"]
    for idx, artifact in enumerate(artifacts, start=1):
        lines.append(f"{idx}. {artifact['filename']} ({artifact['suffix']})")
        lines.append(f"   path: {artifact['path']}")
    parts.append("\n".join(lines))

    content_blocks: list[str] = []
    for artifact in artifacts:
        body = truncate_text(artifact.get("content") or "", MAX_GENERATED_CONTENT_CHARS_PER_FILE)
        if not body:
            continue
        content_blocks.append(f"### Extracted content: {artifact['filename']}\n{body}")
    if content_blocks:
        parts.append("GENERATED FILE CONTENTS:\n\n" + "\n\n".join(content_blocks))
    return "\n\n".join(parts)


def build_output_payload(narrative: str, artifacts: list[dict[str, Any]]) -> dict[str, Any]:
    return {
        "response_text": format_artifact_output(narrative, artifacts),
        "narrative_response": narrative.strip(),
        "generated_files": [
            {
                "filename": artifact["filename"],
                "path": artifact["path"],
                "attachment": make_file_attachment(artifact["path"], artifact["filename"]),
            }
            for artifact in artifacts
        ],
    }


def output_to_scoring_text(output: Any) -> str:
    if isinstance(output, dict):
        response_text = output.get("response_text")
        if isinstance(response_text, str):
            return response_text
        narrative = output.get("narrative_response")
        if isinstance(narrative, str):
            return narrative
    return str(output)


def normalize_rubric_score(earned_points: float, rubric_items: list[dict[str, Any]]) -> float:
    min_points = float(sum(float(item["score"]) for item in rubric_items if float(item["score"]) < 0))
    max_points = float(sum(float(item["score"]) for item in rubric_items if float(item["score"]) > 0))
    score_range = max_points - min_points
    if score_range <= 0:
        return 0.0
    normalized = (earned_points - min_points) / score_range
    return max(0.0, min(1.0, normalized))


def normalize_chat_completion_kwargs(kwargs: dict[str, Any]) -> dict[str, Any]:
    normalized = dict(kwargs)
    model = str(normalized.get("model") or "")
    if model.startswith("gpt-5") and "max_tokens" in normalized and "max_completion_tokens" not in normalized:
        normalized["max_completion_tokens"] = normalized.pop("max_tokens")
    return normalized


def _retry_delay_from_error(err: RateLimitError, fallback_delay: float) -> float:
    response = getattr(err, "response", None)
    if response is not None:
        headers = getattr(response, "headers", {}) or {}
        retry_after = headers.get("retry-after-ms") or headers.get("x-ratelimit-reset-requests-ms")
        if retry_after:
            try:
                return max(float(retry_after) / 1000.0, fallback_delay)
            except ValueError:
                pass
        retry_after = headers.get("retry-after")
        if retry_after:
            try:
                return max(float(retry_after), fallback_delay)
            except ValueError:
                pass

    message = str(err)
    match = re.search(r"try again in ([0-9.]+)ms", message, re.IGNORECASE)
    if match:
        return max(float(match.group(1)) / 1000.0, fallback_delay)
    match = re.search(r"try again in ([0-9.]+)s", message, re.IGNORECASE)
    if match:
        return max(float(match.group(1)), fallback_delay)
    return fallback_delay


def create_chat_completion_with_retries(
    client: OpenAI,
    **kwargs: Any,
):
    delay = INITIAL_OPENAI_BACKOFF_SEC
    last_error: Exception | None = None
    request_kwargs = normalize_chat_completion_kwargs(kwargs)
    for attempt in range(MAX_OPENAI_RETRIES):
        try:
            return client.chat.completions.create(**request_kwargs)
        except RateLimitError as err:
            last_error = err
            if attempt == MAX_OPENAI_RETRIES - 1:
                break
            delay = _retry_delay_from_error(err, delay)
            time.sleep(delay)
            delay = min(delay * 2, 30.0)
        except BadRequestError as err:
            last_error = err
            error_body = getattr(err, "body", {}) or {}
            error_details = error_body.get("error", {}) if isinstance(error_body, dict) else {}
            unsupported_param = error_details.get("param")
            message = str(error_details.get("message") or err)
            if (
                unsupported_param == "max_tokens"
                and "max_completion_tokens" in message
                and "max_tokens" in request_kwargs
                and "max_completion_tokens" not in request_kwargs
            ):
                request_kwargs["max_completion_tokens"] = request_kwargs.pop("max_tokens")
                continue
            raise
    assert last_error is not None
    raise last_error
